import streamlit as st
import os
import json
import tempfile
import base64
from pptx import Presentation

# ==========================================
# CONFIGURATION ET SAUVEGARDE
# ==========================================
st.set_page_config(page_title="Générateur de CV PowerPoint", layout="wide")
DATA_FILE = "cv_data.json"
PROFILE_FIELDS = ['prenom', 'nom', 'poste', 'experience_ans', 'resume', 'expertise', 'positionnement']

def load_data():
    if os.path.exists(DATA_FILE):
        try:
            with open(DATA_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            pass
    return None

def save_data(silent=True):
    data = {
        'profil': st.session_state.profil,
        'competences': st.session_state.competences,
        'diplomes': st.session_state.diplomes,
        'certifications': st.session_state.certifications,
        'experiences': st.session_state.experiences
    }
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)
    if not silent:
        st.sidebar.success("✅ Données sauvegardées localement !")

# ==========================================
# INITIALISATION ROBUSTE DE LA MÉMOIRE
# ==========================================
if 'initialized' not in st.session_state:
    saved = load_data()
    if saved:
        st.session_state.profil = saved.get('profil', {})
        st.session_state.competences = saved.get('competences', [])
        st.session_state.diplomes = saved.get('diplomes', [])
        st.session_state.certifications = saved.get('certifications', [])
        st.session_state.experiences = saved.get('experiences', [])
    else:
        st.session_state.profil = {}
        st.session_state.competences = []
        st.session_state.diplomes = []
        st.session_state.certifications = []
        st.session_state.experiences = []
        
    for f in PROFILE_FIELDS:
        st.session_state[f] = st.session_state.profil.get(f, "")
        
    st.session_state.initialized = True

def sync_profil():
    for f in PROFILE_FIELDS:
        st.session_state.profil[f] = st.session_state[f]

with st.sidebar:
    st.title("💾 Stockage")
    st.success("🟢 **Sauvegarde automatique activée.** \nToutes vos modifications sont enregistrées en temps réel.")

# ==========================================
# FONCTIONS PPTX
# ==========================================
def replace_profile_picture(prs, slide, new_image_path):
    target_shape = None
    for shape in slide.shapes:
        if getattr(shape, "shape_type", None) == 13: # 13 = Image
            if shape.width < prs.slide_width * 0.5:
                target_shape = shape
                break
                
    if target_shape:
        left, top = target_shape.left, target_shape.top
        width, height = target_shape.width, target_shape.height
        sp = target_shape._element
        sp.getparent().remove(sp)
        slide.shapes.add_picture(new_image_path, left, top, width, height)

def generate_pptx_cv(template_path, data):
    prs = Presentation(template_path)
    p = data['profil']
    
    # 0. Remplacement de la photo
    if p.get('photo_b64'):
        img_data = base64.b64decode(p['photo_b64'])
        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as tmp_img:
            tmp_img.write(img_data)
            tmp_img_path = tmp_img.name
        
        replace_profile_picture(prs, prs.slides[0], tmp_img_path)
        try: os.unlink(tmp_img_path)
        except: pass

    exp_str = f"+{p.get('experience_ans','')} années d'expérience" if str(p.get('experience_ans','')).isdigit() else p.get('experience_ans','')
    
    standard_replacements = {
        "[PRÉNOM]": p.get('prenom', ''),
        "[NOM]": p.get('nom', '').upper(),
        "[POSTE / FONCTION]": p.get('poste', ''),
        "[X années d’expérience]": exp_str,
        "[X années d'expérience]": exp_str,
        "[RÉSUMÉ PROFESSIONNEL — 2 à 4 lignes]": p.get('resume', ''),
        "[EXPERTISE / VALEUR AJOUTÉE — 2 à 4 lignes]": p.get('expertise', ''),
        "[POSITIONNEMENT / DOMAINES D’INTERVENTION — 2 à 4 lignes]": p.get('positionnement', '')
    }

    for i in range(4):
        val = data['competences'][i] if i < len(data['competences']) else ""
        standard_replacements[f"#[COMPÉTENCE {i+1}]"] = f"#{val}" if val else ""

    for i in range(2):
        val = f"{data['diplomes'][i]['ecole']} - {data['diplomes'][i]['titre']}" if i < len(data['diplomes']) else ""
        standard_replacements[f"[DIPLÔME / FORMATION {i+1}]"] = val

    for i in range(4):
        val = f"{data['certifications'][i]['nom']}" if i < len(data['certifications']) else ""
        standard_replacements[f"[CERTIFICATION {i+1}]"] = val
        if i == 2: standard_replacements["[CERTIFICATION / LANGUE]"] = val

    for i in range(4):
        idx = i + 1
        if i < len(data['experiences']):
            exp = data['experiences'][i]
            standard_replacements[f"[ENTREPRISE {idx}]"] = exp['entreprise']
            standard_replacements[f"[POSTE {idx}]"] = exp['poste']
            standard_replacements[f"[CONTEXTE / MISSION {idx} — 1 à 2 lignes]"] = exp['contexte']
            standard_replacements[f"[CONTEXTE / MISSION {idx} — 1 à 3 lignes]"] = exp['contexte']
            for j in range(3):
                real_val = exp['realisations'][j] if j < len(exp['realisations']) else ""
                standard_replacements[f"[RÉALISATION {idx}.{j+1}]"] = real_val
        else:
            standard_replacements[f"[ENTREPRISE {idx}]"] = ""
            standard_replacements[f"[POSTE {idx}]"] = ""
            standard_replacements[f"[CONTEXTE / MISSION {idx} — 1 à 2 lignes]"] = ""
            standard_replacements[f"[CONTEXTE / MISSION {idx} — 1 à 3 lignes]"] = ""
            for j in range(3):
                standard_replacements[f"[RÉALISATION {idx}.{j+1}]"] = ""

    paragraphs_with_y = []
    for slide in prs.slides:
        for shape in slide.shapes:
            def extract_paras(shp, current_top):
                if getattr(shp, "shape_type", None) == 6:
                    for c in shp.shapes:
                        extract_paras(c, current_top + getattr(c, "top", 0))
                if getattr(shp, "has_text_frame", False):
                    for paragraph in shp.text_frame.paragraphs:
                        paragraphs_with_y.append((current_top, paragraph))
                if getattr(shp, "has_table", False):
                    for row in shp.table.rows:
                        for cell in row.cells:
                            if getattr(cell, "text_frame", None):
                                for paragraph in cell.text_frame.paragraphs:
                                    paragraphs_with_y.append((current_top, paragraph))
            extract_paras(shape, getattr(shape, "top", 0))
            
    paragraphs_with_y.sort(key=lambda x: x[0])
    
    meta_counter = 0
    paragraphs_to_delete = []
    
    for _, paragraph in paragraphs_with_y:
        original_text = "".join(run.text for run in paragraph.runs)
        if not original_text.strip(): continue 
            
        meta_tag = "[TYPE / CHARGE / DURÉE / LOCALISATION]"
        if meta_tag in original_text:
            if meta_counter < len(data['experiences']):
                exp = data['experiences'][meta_counter]
                parts = [p for p in [exp.get('type'), exp.get('charge'), exp.get('duree'), exp.get('localisation')] if p]
                meta_val = " – ".join(parts[:2]) + (f" ({parts[2]})" if len(parts)>2 else "") + (f", {parts[3]}" if len(parts)>3 else "")
            else:
                meta_val = ""
            standard_replacements[meta_tag] = meta_val
            meta_counter += 1

        for key, val in standard_replacements.items():
            if key in original_text:
                replaced_in_run = False
                for run in paragraph.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, str(val))
                        replaced_in_run = True
                
                if not replaced_in_run:
                    full_text = "".join(r.text for r in paragraph.runs)
                    new_text = full_text.replace(key, str(val))
                    if paragraph.runs:
                        paragraph.runs[0].text = new_text
                        for r in paragraph.runs[1:]: r.text = ""
                original_text = "".join(run.text for run in paragraph.runs)

        final_text = "".join(run.text for run in paragraph.runs)
        if not final_text.strip(" -–—,.|•\t\n\r"):
            paragraphs_to_delete.append(paragraph)

    for p in paragraphs_to_delete:
        p_element = p._p
        parent = p_element.getparent()
        if parent is not None:
            parent.remove(p_element)
                
    tmp_pptx = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    prs.save(tmp_pptx.name)
    tmp_pptx.close()
    return tmp_pptx.name

# ==========================================
# INTERFACE UTILISATEUR
# ==========================================
st.title("💼 Générateur de CV Avancé")

tab1, tab2, tab3 = st.tabs(["1. Mon Profil", "2. Parcours & Expériences", "3. Génération & Export"])

# --- ONGLET 1 : PROFIL ---
with tab1:
    st.header("Informations générales")
    
    col_img, col_info = st.columns([1, 3])
    with col_img:
        st.write("📷 **Photo de profil**")
        if st.session_state.profil.get('photo_b64'):
            st.image(base64.b64decode(st.session_state.profil['photo_b64']), width=150)
        
        uploaded_file = st.file_uploader("Modifier la photo", type=["jpg", "jpeg", "png"], label_visibility="collapsed")
        if uploaded_file is not None:
            st.session_state.profil['photo_b64'] = base64.b64encode(uploaded_file.read()).decode()
            st.rerun()

    with col_info:
        c1, c2 = st.columns(2)
        c1.text_input("Prénom", key="prenom", on_change=sync_profil)
        c2.text_input("Nom", key="nom", on_change=sync_profil)
        c1.text_input("Poste / Fonction", key="poste", placeholder="Ex: Product Owner", on_change=sync_profil)
        c2.text_input("Années d'expérience", key="experience_ans", placeholder="Ex: 3", on_change=sync_profil)
    
    st.subheader("Textes de présentation")
    st.text_area("Résumé (2-4 lignes)", key="resume", help="Ex: Issu du Programme...", height=80, on_change=sync_profil)
    st.text_area("Expertise / Valeur ajoutée", key="expertise", help="Ex: En pilotant des projets...", height=80, on_change=sync_profil)
    st.text_area("Domaines d'intervention", key="positionnement", help="Ex: Dan intervient à la fois...", height=80, on_change=sync_profil)

# --- ONGLET 2 : AJOUT DE DONNÉES ---
with tab2:
    st.header("Alimenter la base de données")
    
    c1, c2 = st.columns(2)
    
    with c1:
        with st.form("form_comp", clear_on_submit=True):
            st.subheader("Tags / Compétences")
            new_comp = st.text_input("Nouveau Tag (sans le #)", placeholder="Ex: discovery, erp, data...")
            if st.form_submit_button("➕ Ajouter tag") and new_comp:
                if new_comp not in st.session_state.competences: 
                    st.session_state.competences.append(new_comp)
        
        if st.session_state.competences:
            st.success("✅ **Enregistré :** " + ", ".join(st.session_state.competences))
            if st.button("🗑️ Vider tags", key="clr_tag"):
                st.session_state.competences = []
                st.rerun()
                    
        with st.form("form_dip", clear_on_submit=True):
            st.subheader("Diplômes")
            ecole = st.text_input("École", placeholder="Ex: ESSEC Business School")
            titre = st.text_input("Titre du diplôme", placeholder="Ex: Master in Management")
            if st.form_submit_button("➕ Ajouter diplôme") and ecole:
                st.session_state.diplomes.append({"ecole": ecole, "titre": titre})
                
        if st.session_state.diplomes:
            st.info("🎓 **Diplômes enregistrés :**\n" + "\n".join([f"- {d['ecole']} ({d['titre']})" for d in st.session_state.diplomes]))
            if st.button("🗑️ Vider diplômes", key="clr_dip"):
                st.session_state.diplomes = []
                st.rerun()
                
        with st.form("form_cert", clear_on_submit=True):
            st.subheader("Certifications & Langues")
            nom_cert = st.text_input("Nom de la certification", placeholder="Ex: Certification PSPO")
            if st.form_submit_button("➕ Ajouter certification") and nom_cert:
                st.session_state.certifications.append({"nom": nom_cert})
                
        if st.session_state.certifications:
            st.info("🏅 **Certifications enregistrées :**\n" + "\n".join([f"- {c['nom']}" for c in st.session_state.certifications]))
            if st.button("🗑️ Vider certifications", key="clr_cert"):
                st.session_state.certifications = []
                st.rerun()

    with c2:
        with st.form("form_exp", clear_on_submit=True):
            st.subheader("Expérience Professionnelle")
            ent = st.text_input("Entreprise", placeholder="Ex: EDMOND DE ROTHSCHILD")
            poste = st.text_input("Poste", placeholder="Ex: Chef de projet")
            ctx = st.text_area("Contexte / Mission", placeholder="Ex: Accompagnement du groupe...")
            
            c_meta1, c_meta2 = st.columns(2)
            meta_type = c_meta1.text_input("Type de contrat", placeholder="Ex: Forfait")
            meta_charge = c_meta2.text_input("Charge", placeholder="Ex: temps plein")
            meta_duree = c_meta1.text_input("Durée", placeholder="Ex: 2 mois")
            meta_loc = c_meta2.text_input("Lieu", placeholder="Ex: Paris")
            
            reals = st.text_area("Réalisations (Listez 3 points maximum, un par ligne)")
            
            if st.form_submit_button("➕ Enregistrer l'expérience") and ent:
                real_list = [r.strip() for r in reals.split('\n') if r.strip()][:3]
                st.session_state.experiences.append({
                    "entreprise": ent, "poste": poste, "contexte": ctx,
                    "type": meta_type, "charge": meta_charge, "duree": meta_duree, "localisation": meta_loc,
                    "realisations": real_list
                })
                
        if st.session_state.experiences:
            st.markdown("### 💼 Expériences enregistrées dans la base")
            for i, e in enumerate(st.session_state.experiences):
                st.success(f"**{e['entreprise']}** - {e['poste']}")
            if st.button("🗑️ Vider les expériences", key="clr_exp"):
                st.session_state.experiences = []
                st.rerun()

# --- ONGLET 3 : GÉNÉRATION ---
with tab3:
    st.header("Génération du document")
    st.write("Sélectionnez les éléments à inclure pour **ce CV spécifiquement**.")
    
    col_sel1, col_sel2 = st.columns(2)
    with col_sel1:
        selected_comps = st.multiselect("Compétences (Max 4 affichées)", st.session_state.competences, default=st.session_state.competences[:4])
        dip_opts = [f"{d['ecole']} - {d['titre']}" for d in st.session_state.diplomes]
        sel_dip = st.multiselect("Diplômes (Max 2 affichés)", dip_opts, default=dip_opts[:2])
        selected_diplomes = [d for d in st.session_state.diplomes if f"{d['ecole']} - {d['titre']}" in sel_dip]
    
    with col_sel2:
        cert_opts = [c['nom'] for c in st.session_state.certifications]
        sel_cert = st.multiselect("Certifications (Max 4 affichées)", cert_opts, default=cert_opts[:4])
        selected_certifications = [c for c in st.session_state.certifications if c['nom'] in sel_cert]
        
        exp_opts = [f"{e['entreprise']} - {e['poste']}" for e in st.session_state.experiences]
        sel_exp = st.multiselect("Expériences (Max 4 affichées)", exp_opts, default=exp_opts[:4])
        selected_experiences = [e for e in st.session_state.experiences if f"{e['entreprise']} - {e['poste']}" in sel_exp]

    template_file = "CV_Template_reutilisable.pptx"
    
    st.write("---")
    if not os.path.exists(template_file):
        st.error(f"❌ Fichier introuvable : `{template_file}`. Veuillez le placer dans le même dossier.")
    else:
        if st.button("🚀 GÉNÉRER MON CV", type="primary", use_container_width=True):
            data_to_render = {
                "profil": st.session_state.profil,
                "competences": selected_comps,
                "diplomes": selected_diplomes,
                "certifications": selected_certifications,
                "experiences": selected_experiences
            }
            
            with st.spinner("Création du fichier PowerPoint..."):
                try:
                    pptx_path = generate_pptx_cv(template_file, data_to_render)
                    
                    with open(pptx_path, "rb") as file:
                        st.download_button(
                            label="📥 Télécharger le CV (PPTX)", 
                            data=file, 
                            file_name=f"CV_{st.session_state.profil.get('nom', 'Genere')}.pptx", 
                            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", 
                            use_container_width=True
                        )
                    st.info("💡 **Pour obtenir un PDF :** Ouvrez le fichier téléchargé avec PowerPoint, puis faites `Fichier > Enregistrer sous > Format PDF`.")
                            
                except Exception as e:
                    st.error(f"❌ Erreur lors du remplacement : {e}")

# ==========================================
# SAUVEGARDE AUTOMATIQUE EN FIN DE SCRIPT
# ==========================================
save_data(silent=True)